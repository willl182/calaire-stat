"""Números docentes simulados del capítulo M00; no son datos de CALAIRE."""

from collections import Counter
from pathlib import Path
import subprocess
import tempfile

from openpyxl import Workbook, load_workbook

# B: dos tablas de origen con una clave de pareja explícita.
referencias = {
    "P01": ("L0", 0.0),
    "P02": ("L1", 50.0),
    "P03": ("L1", 50.0),
    "P04": ("L2", 100.0),
    "P05": ("L3", 150.0),
}
indicaciones = {
    "P01": -0.2,  # observación del cero: se conserva con su signo
    "P02": 50.4,
    # No hay indicación P03: el faltante no se sustituye por la lectura siguiente.
    "P04": 100.8,
    "P05": 150.7,
}

pares = []
for id_pareja, (nivel, x) in referencias.items():
    y = indicaciones.get(id_pareja)
    e = None if y is None else y - x
    pares.append((id_pareja, nivel, x, y, e))

validos = [fila for fila in pares if fila[4] is not None]
media_e = sum(fila[4] for fila in validos) / len(validos)
lecturas_sin_clave = list(indicaciones.values())
pareo_posicional_incorrecto = lecturas_sin_clave[2] - referencias["P03"][1]

print("B: parejas (id, nivel, x, y, e)")
for fila in pares:
    print(fila)
print(f"B: referencias={len(referencias)}, indicaciones={len(indicaciones)}, "
      f"parejas válidas={len(validos)}, faltantes={len(referencias) - len(validos)}")
print(f"B: media de errores válidos={media_e:.3f} nmol/mol")
print(f"B: falso error al emparejar por posición para P03="
      f"{pareo_posicional_incorrecto:.1f} nmol/mol")

# I: ejemplo balanceado simulado; cada fila representa un ciclo observado.
registros = []
for campania in ("C1", "C2"):
    for analista in ("A1", "A2"):
        for dia in (1, 2):
            for ciclo in (1, 2):
                registros.append((campania, analista, dia, ciclo))

por_campania = Counter(r[0] for r in registros)
por_analista = Counter(r[1] for r in registros)
por_campania_analista = Counter((r[0], r[1]) for r in registros)
por_campania_analista_dia = Counter((r[0], r[1], r[2]) for r in registros)
assert len(por_campania_analista_dia) == 8
assert set(por_campania_analista_dia.values()) == {2}
print(f"I: total={len(registros)}; por campaña={dict(sorted(por_campania.items()))}; "
      f"por analista={dict(sorted(por_analista.items()))}")
print(f"I: conteos campaña×analista={dict(sorted(por_campania_analista.items()))}")
print(f"I: conteos campaña×analista×día={dict(sorted(por_campania_analista_dia.items()))}")

# I2: ejemplo simulado de transformación ancho a largo con un dato faltante.
tabla_ancha = [
    ("C1", "A1", 1, 10.0, 11.0),
    ("C1", "A2", 1, 12.0, None),
]
registros_largos = [
    (campania, analista, dia, ciclo, valor)
    for campania, analista, dia, lectura_1, lectura_2 in tabla_ancha
    for ciclo, valor in ((1, lectura_1), (2, lectura_2))
    if valor is not None
]
valores_no_faltantes = sum(
    valor is not None for fila in tabla_ancha for valor in fila[3:]
)
faltantes_largos = [
    (campania, analista, dia, ciclo, "faltante")
    for campania, analista, dia, lectura_1, lectura_2 in tabla_ancha
    for ciclo, valor in ((1, lectura_1), (2, lectura_2))
    if valor is None
]
assert len(registros_largos) == valores_no_faltantes == 3
assert faltantes_largos == [("C1", "A2", 1, 2, "faltante")]
assert len(registros_largos) + len(faltantes_largos) == 2 * len(tabla_ancha)
print(f"I2: estados faltantes conservados={faltantes_largos}")
print(f"I2: filas anchas={len(tabla_ancha)}; celdas de lectura no faltantes="
      f"{valores_no_faltantes}; registros largos={len(registros_largos)}")
print(f"I2: registros largos={registros_largos}")

# A: comparar Python con LibreOffice Calc; x e y se guardan como entradas y
# Calc calcula en cada fila el error con la fórmula =C2-B2.
x_tool = [50.0, 100.0, 150.0]
y_tool = [50.4, 100.8, 150.7]
e_tool_python = [y - x for x, y in zip(x_tool, y_tool)]
with tempfile.TemporaryDirectory(dir=Path(__file__).parent) as tmp:
    root = Path(tmp)
    source_dir, output_dir, profile_dir = (root / "source", root / "output", root / "profile")
    source_dir.mkdir()
    output_dir.mkdir()
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(("pareja", "referencia_x", "indicacion_y", "error_e"))
    for index, (x, y) in enumerate(zip(x_tool, y_tool), start=1):
        row = index + 1
        sheet.append((f"P{index:02d}", x, y, f"=C{row}-B{row}"))
    source = source_dir / "comparacion.xlsx"
    workbook.save(source)
    subprocess.run(
        ["libreoffice", f"-env:UserInstallation={profile_dir.as_uri()}", "--headless",
         "--convert-to", "xlsx", "--outdir", str(output_dir), str(source)],
        check=True,
        capture_output=True,
        text=True,
        timeout=30,
    )
    calculated = load_workbook(output_dir / source.name, data_only=True)
    sheet_values = calculated.active
    e_tool_calc = [sheet_values.cell(row, 4).value for row in range(2, 5)]

diferencia_maxima = max(abs(a - b) for a, b in zip(e_tool_python, e_tool_calc))
tolerancia = 1e-12
assert len(e_tool_calc) == len(e_tool_python) == 3
assert diferencia_maxima <= tolerancia, "Python y Calc exceden la tolerancia declarada"
print(f"A: errores Python (redondeados a 3 decimales)="
      f"{[round(v, 3) for v in e_tool_python]}")
print(f"A: errores LibreOffice Calc (redondeados a 3 decimales)="
      f"{[round(v, 3) for v in e_tool_calc]}")
print(f"A: igualdad binaria exacta={e_tool_python == e_tool_calc}; "
      f"diferencia máxima={diferencia_maxima:.3e} nmol/mol; "
      f"acuerdo con tolerancia {tolerancia:.0e} nmol/mol="
      f"{diferencia_maxima <= tolerancia}")

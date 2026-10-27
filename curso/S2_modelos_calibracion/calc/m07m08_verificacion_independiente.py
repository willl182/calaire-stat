"""Contrastes independientes: fórmulas centradas y GGMR perfilado (sin importar scripts)."""
from pathlib import Path
from math import sqrt
import numpy as np
from scipy.optimize import minimize_scalar
from scipy.stats import t, f

x=np.arange(0.,501.,100.); y=np.array([.2,100.7,200.1,301.,400.4,501.2])
sxx=sum((x-x.mean())**2)
b=sum((x-x.mean())*(y-y.mean()))/sxx; a=y.mean()-b*x.mean()
s2=sum((y-a-b*x)**2)/4
xh=(250.5-a)/b
vp=s2*(1/6+(xh-x.mean())**2/sxx)
u=sqrt((.2**2+vp)/b**2)
nu_eff=(.2**2+vp)**2/(vp**2/4)
print('OLS centrado',a,b,sqrt(s2),'inversa',xh,u,nu_eff,t.ppf(.975,nu_eff)*u)
assert abs(u-.25644562)<5e-9
xr=np.repeat([0.,100.,200.,300.],3)
yr=np.array([.1,.2,.3,100.2,100.1,100.3,200.,200.1,200.2,301.,301.1,301.2])
br=sum((xr-xr.mean())*(yr-yr.mean()))/sum((xr-xr.mean())**2)
ar=yr.mean()-br*xr.mean()
means=yr.reshape(4,3).mean(1)
pe=sum(((yr.reshape(4,3)-means[:,None])**2).ravel())
lof=sum((yr-ar-br*xr)**2)-pe
print('LOF por diferencia SSE-PE',lof,pe,(lof/2)/(pe/8),f.sf((lof/2)/(pe/8),2,8))

# Se eliminan las abscisas y el intercepto para minimizar una función escalar de b.
def profile(x,y,vx,vy):
    one=np.ones(len(x))
    def evaluate(b, full=False):
        v=vy+b*b*vx
        z=y-b*x
        a=(one@np.linalg.solve(v,z))/(one@np.linalg.solve(v,one))
        r=z-a
        q=r@np.linalg.solve(v,r)
        return (a,b,q) if full else q
    fit=minimize_scalar(evaluate,bounds=(.9,1.1),method='bounded',options={'xatol':1e-14})
    assert fit.success
    return evaluate(fit.x,True)
yg=np.array([.3,100.4,200.7,300.5,401.,500.9])
vx=.25**2*np.eye(6)+.12**2*np.ones((6,6)); vy=.4**2*np.eye(6)
pg=profile(x,yg,vx,vy); print('GGMR perfilado simulado',pg)
assert abs(pg[2]-.40890065)<5e-9
xp=np.array([50.4,99.,149.9,200.4,248.5,299.7,349.1])
yp=np.array([52.3,97.8,149.7,200.1,250.4,300.9,349.2])
# Reconstrucción desde el modelo generador D.4, independiente de transcribir Ux.
A=np.array([[1,0,0],[0,1,0],[1,1,0],[0,0,1],[1,0,1],[0,1,1],[1,1,1]])
vxp=A@np.diag([.5**2,1.,1.])@A.T+.5**2*np.eye(7)
vyp=4*np.eye(7)+np.ones((7,7))
pp=profile(xp,yp,vxp,vyp); print('ISO GGMR perfilado',pp)
assert np.allclose(pp,[.34240092,1.00123076,1.77184745],atol=3e-6,rtol=0)
cy=1/1.002
variance=cy**2*(.24**2/8+.10**2/(12*8)+.15**2)+ (100/1.002*.002)**2
uc=sqrt(variance); df=variance**2/((cy*.24/sqrt(8))**4/7)
print('M08 lineal independiente',uc,df,t.ppf(.975,df)*uc)
assert abs(uc-.2636774)<5e-8
cov=100*102*.003**2
ud=sqrt(.20**2+.25**2+((102-100)*.003)**2)
root=(.0064+sqrt(.0064**2+4*.000128*.3264))/(2*.000128)
print('M08 diferencia por componentes latentes',cov,ud,'cruce fórmula cuadrática',root)
assert abs(ud-.3202124295)<1e-9
assert 81.3<root<81.4
print('M08 individual/promedio/futuro',sqrt(.30**2+.40**2+.10**2),sqrt(.14),sqrt(.14+.25**2+.18**2))

# Comprobación de datos propios sin modificar el libro ni ejecutar sus macros.
from openpyxl import load_workbook
r=Path(__file__).resolve().parents[3]
p=r/'docs/F-PSMT-01 Certificado de Calibracion_G1-sn59756-324_TBH.xlsm'
wb=load_workbook(p,data_only=True,read_only=True)
ws=wb['BD_DIAS']
print('Libro propio BD_DIAS dimensiones',ws.max_row,ws.max_column)
for row in ws.iter_rows():
    vals=[(c.coordinate,c.value) for c in row if c.value is not None]
    if any(isinstance(v,str) and ('Datos por' in v or 'Total Datos' in v) for _,v in vals):
        print('Conteos propios',vals)
wb.close()
# Aserciones adicionales: no basta con imprimir resultados para verificarlos.
assert np.allclose([a,b,sqrt(s2),xh], [.242857142857,1.001428571429,.394606494,.24990014265e3], atol=1e-8, rtol=0)
assert np.allclose([lof,pe,(lof/2)/(pe/8)], [.966,.080,48.3], atol=1e-9, rtol=0)
assert abs(df-657.9543111)<1e-5
assert abs(t.ppf(.975,df)*uc-.5177507)<5e-8
assert abs(sqrt(.28**2+.35**2+.42**2)-.6142475074)<1e-9
# Tabla Kragten por identidades algebraicas, sin reutilizar bucle del script.
uy=sqrt(.24**2/8+.10**2/96)
effects=np.array([uy/1.002,-.15/1.002,-100*.002/(1.002+.002)])
assert abs(np.linalg.norm(effects)-.2633766)<5e-8
assert abs((2*sqrt(.28**2+(.006*root)**2))-(.8+.004*root))<1e-12
# En todo el intervalo el polinomio tiene una sola raíz positiva.
other=(.0064-sqrt(.0064**2+4*.000128*.3264))/(2*.000128)
assert other<0<root<500
for chapter in ('M07','M08'):
    text=(r/'curso/S2_modelos_calibracion'/f'{chapter}_libreto.md').read_text()
    assert text.count('**[PRÁCTICA — fase de ejercicios]**')==3
print('Todas las aserciones independientes superadas')
